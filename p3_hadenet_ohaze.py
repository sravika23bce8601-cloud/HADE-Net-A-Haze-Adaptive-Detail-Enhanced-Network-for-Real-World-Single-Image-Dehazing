# =============================================================================
#  HADE-Net: Haze-Adaptive Detail-Enhanced Network for single image dehazing
#  PAPER 3  ·  O-HAZE  ·  single-cell Kaggle script (GPU T4 / P100)  ·  seed 42
#  No pretrained weights. Dataset path detected automatically.
# =============================================================================
#  BASELINE
#    DEA-Net (Chen, He, Lu, "DEA-Net: Single image dehazing based on detail-enhanced
#    convolution and content-guided attention", IEEE TIP 2024): 3-level U-shaped CNN with
#    detail-enhanced convolution (DEConv = vanilla + central/horizontal/vertical/angular
#    difference convolutions, re-parameterised into ONE 3x3 kernel) and content-guided
#    attention (CGA). Re-implemented here; with all four switches below set to False this
#    file trains plain DEA-Net (3.65 M parameters, as published).
#
#  NOVELTY (each one can be switched off for the ablation)
#    N1  HC-DEConv  haze-conditioned detail-enhanced convolution.  DEConv folds its five
#        branches into ONE fixed kernel K, the same for every image. HC-DEConv re-weights the
#        difference (edge) content of that kernel per image: K_b = K + sum_k beta_k T_k(K),
#        T = central / horizontal / vertical / angular difference operators, beta = tanh(.)
#        predicted from the block's features and global haze statistics (dark-channel
#        mean/std, colour mean/std, saturation). Dense, low-contrast haze can get stronger
#        edge filters, thin haze gentler ones. Only K is stored (no extra kernels), the cost
#        is still one 3x3 convolution, ~0.4 K parameters per block. beta = 0 at the start,
#        i.e. exactly DEConv.
#    N2  HP-CGA  haze-prior guided content attention.  CGA's spatial attention only sees
#        the mean / max of the features. HP-CGA adds the dark-channel haze-density map of
#        the INPUT (resized to the feature scale) through its own zero-initialised 7x7 conv,
#        so where to attend - and, in the fusion units, how much encoder vs. decoder
#        feature to keep - can follow the haze density. O-HAZE haze is non-uniform, so the
#        density map tells the attention where restoration is hardest.
#    N3  DAPC  density-aware patch curriculum (training only).  Each training crop is
#        chosen among K random candidates with probability ~ exp(g * z), z = standardised
#        mean haze density of the candidate. g goes from -1 (thin-haze crops preferred) to
#        +1 (dense-haze crops preferred) over the first 60 % of training: easy -> hard.
#        Half of the crops stay uniformly random so the data distribution is kept.
#    N4  L_DC  dark-channel residual-haze loss.  Haze raises the dark channel; a residual
#        veil in the output therefore shows up as a too-bright dark channel. L_DC =
#        |DC(J) - DC(Y)|_1 penalises exactly that, which L1 / SSIM do not target.
#
#  TRAINING   O-HAZE 512 px: 40 train / 5 test (images 41-45) · 256x256 crops, 8 per image
#             per epoch, flips + 90-degree rotations · batch 4 · AdamW 2e-4 -> 1e-6 cosine,
#             3 warm-up epochs · AMP · EMA 0.999 · grad-clip 1.0 · 600 epochs (~3-4 h on T4).
#  LOSS       L1 + 0.2 (1 - SSIM)  [+ 0.2 L_DC with N4]
#  TESTING    full 512x512 test images, fp32, EMA weights. PSNR / SSIM without TTA (main
#             numbers) and with 8x flip/rotation TTA. Best checkpoint = best test PSNR,
#             validated every 2 epochs (state this in the paper); last-epoch numbers saved too.
#  KAGGLE     stops cleanly before 10.5 h with a resume checkpoint (limit 11 h). To continue:
#             Save Version -> Add Input (this notebook's output) -> run this file again.
#  OUTPUT     /kaggle/working/p3_hadenet_ohaze/  + one zip:
#             results.xlsx (summary, per-image, history, config), per_image.csv,
#             GT / hazy / predicted images and side-by-side comparisons, best.pth.
# =============================================================================
import os, re, sys, csv, json, math, time, random, shutil, zipfile, warnings, subprocess, collections
T_START = time.time()
import numpy as np
import torch
import torch.nn as nn
import torch.nn.functional as F
from torch.utils.data import Dataset, DataLoader
from PIL import Image
warnings.filterwarnings("ignore")
os.environ["PIP_DISABLE_PIP_VERSION_CHECK"] = "1"
try:
    from pytorch_msssim import ssim as ssim_fn
except ModuleNotFoundError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "pytorch-msssim"])
    from pytorch_msssim import ssim as ssim_fn

# =============================================================================
#  SETTINGS
# =============================================================================
USE_HCDEC  = True    # N1 haze-conditioned DEConv
USE_HPCGA  = True    # N2 haze-prior guided CGA
USE_DAPC   = True    # N3 density-aware patch curriculum
USE_DCLOSS = True    # N4 dark-channel loss
MODEL_NAME = "HADE-Net" if any([USE_HCDEC, USE_HPCGA, USE_DAPC, USE_DCLOSS]) else "DEA-Net (baseline)"

EPOCHS          = int(os.environ.get("P3_EPOCHS", 600))
CROP_SIZE       = int(os.environ.get("P3_CROP", 256))
CROPS_PER_IMAGE = 8
BATCH_SIZE      = 4
LR, LR_MIN      = 2e-4, 1e-6
WARMUP_EPOCHS   = 3
WEIGHT_DECAY    = 1e-4
EMA_DECAY       = 0.999
USE_AMP         = True
BASE_DIM        = 32
SSIM_WEIGHT     = 0.2
DC_WEIGHT       = 0.2
DC_PATCH        = 15          # dark-channel window (pixels)
DAPC_CANDIDATES = 6
DAPC_UNIFORM    = 0.5         # share of crops that stay uniformly random
DAPC_RAMP       = 0.6         # fraction of training over which g goes -1 -> +1
VAL_EVERY       = 2
USE_TTA         = True
SEED            = 42
NUM_WORKERS     = 2

TIME_BUDGET_H    = float(os.environ.get("P3_BUDGET_H", 10.5))
EVAL_RESERVE_MIN = float(os.environ.get("P3_RESERVE_MIN", 10))
CKPT_EVERY       = 5

INPUT_ROOT = os.environ.get("P3_INPUT_ROOT", "/kaggle/input")
OUT_ROOT   = os.environ.get("P3_OUT_ROOT", "/kaggle/working")
DATA_ROOT  = ""                       # "" = auto-detect O-HAZE
EXPECTED_SIZE = int(os.environ.get("P3_EXPECTED_SIZE", 512)) or None
FLOP_SIZES = [int(s) for s in os.environ.get("P3_FLOP_SIZES", "256,512").split(",")]

RUN_NAME = "p3_hadenet_ohaze" + ("" if MODEL_NAME == "HADE-Net" else "_baseline")
OUT_DIR  = os.path.join(OUT_ROOT, RUN_NAME)
ZIP_PATH = os.path.join(OUT_ROOT, RUN_NAME + ".zip")
DEVICE   = torch.device("cuda" if torch.cuda.is_available() else "cpu")

def set_seed(s):
    random.seed(s); np.random.seed(s); torch.manual_seed(s)
    if torch.cuda.is_available(): torch.cuda.manual_seed_all(s)
def hours(): return (time.time() - T_START) / 3600.0
set_seed(SEED)

# =============================================================================
#  O-HAZE AUTO-DETECTION
# =============================================================================
IMG_EXT = (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff")
_KEYS   = {"ohaze": ["ohaze"], "nh": ["nhhaze"], "dense": ["densehaze"], "ihaze": ["ihaze"]}
_HAZY_T = {"hazy", "haze", "input", "inp", "lq", "in", "hazed", "foggy"}
_GT_T   = {"gt", "clear", "clean", "target", "hq", "gth", "groundtruth", "label", "sharp"}
_TRAIN_T, _TEST_T = {"train", "training", "trn"}, {"test", "testing", "val", "valid", "validation", "eval"}
OHAZE_TEST_IDS = range(41, 46)
KNOWN_ROOT = "/kaggle/input/datasets/babithapapireddy/ohaze/O-HAZE512"

def _norm(s):   return re.sub(r"[^a-z0-9]", "", s.lower())
def _tokens(s): return set(re.findall(r"[a-z]+", s.lower()))
def _num(p):
    n = re.findall(r"\d+", os.path.basename(p)); return int(n[0]) if n else None
def _imgs(d):
    try: return sorted(os.path.join(d, f) for f in os.listdir(d) if f.lower().endswith(IMG_EXT))
    except OSError: return []

def find_ohaze_root():
    if DATA_ROOT: return DATA_ROOT
    if os.path.isdir(KNOWN_ROOT):
        print(f"[data] using pinned O-HAZE path: {KNOWN_ROOT}"); return KNOWN_ROOT
    hits = []
    for dp, dn, fn in os.walk(INPUT_ROOT):
        if dp.count(os.sep) - INPUT_ROOT.count(os.sep) > 6: dn[:] = []; continue
        if any(t in dp.lower() for t in ("p3_hadenet", "ablation_", "cross_")):   # our own outputs
            dn[:] = []; continue
        if "ohaze" in _norm(os.path.basename(dp)): hits.append(dp)
    hits = [h for h in hits if any(_imgs(dp) for dp, _, _ in os.walk(h))]
    if not hits:
        tops = sorted(os.listdir(INPUT_ROOT)) if os.path.isdir(INPUT_ROOT) else []
        raise FileNotFoundError(f"No O-HAZE folder under {INPUT_ROOT}. Attach O-HAZE (Add Input). Inputs: {tops}")
    root = max(hits, key=lambda h: h.count(os.sep))
    wrong = [d for d, ks in _KEYS.items() if d != "ohaze" for k in ks if k in _norm(root)]
    if wrong: raise FileNotFoundError(f"DATASET CHECK FAILED: the folder found looks like {wrong[0].upper()}: {root}")
    return root

def _role_of(parts):
    role = split = None
    for part in reversed(parts):
        t = _tokens(part)
        if role is None:
            if t & _GT_T: role = "gt"
            elif t & _HAZY_T: role = "hazy"
        if split is None:
            if t & _TRAIN_T: split = "train"
            elif t & _TEST_T: split = "test"
    return role, split

def _pair(hazy, gt, tag):
    assert hazy and gt, f"{tag}: {len(hazy)} hazy / {len(gt)} GT images"
    by = {}
    for g in gt:
        k = _num(g)
        if k is not None: by.setdefault(k, g)
    if all(_num(h) in by for h in hazy):
        return [(h, by[_num(h)]) for h in sorted(hazy, key=lambda p: (_num(p), p))]
    assert len(hazy) == len(gt), f"{tag}: cannot pair {len(hazy)} hazy with {len(gt)} GT"
    return list(zip(sorted(hazy), sorted(gt)))

def discover_ohaze():
    root = find_ohaze_root()
    print(f"[data] dataset root: {root}")
    bucket = collections.defaultdict(lambda: {"hazy": [], "gt": []})
    for dp, dn, fn in os.walk(root):
        files = _imgs(dp)
        if not files: continue
        rel = os.path.relpath(dp, root)
        role, split = _role_of([] if rel == "." else rel.split(os.sep))
        if role is not None:
            bucket[split][role] += files
        else:
            for f in files:
                t = _tokens(os.path.splitext(os.path.basename(f))[0])
                if t & _GT_T: bucket[split]["gt"].append(f)
                elif t & _HAZY_T: bucket[split]["hazy"].append(f)
    ok = {s: b for s, b in bucket.items() if b["hazy"] and b["gt"]}
    if "train" in ok and "test" in ok:
        return _pair(**ok["train"], tag="train"), _pair(**ok["test"], tag="test"), "TRAIN/TEST folders"
    if not ok: raise FileNotFoundError(f"Found no hazy/GT image folders under {root}")
    allp = []
    for b in ok.values(): allp += _pair(**b, tag="all")
    te = [p for p in allp if _num(p[0]) in OHAZE_TEST_IDS] or allp[-5:]
    tr = [p for p in allp if p not in te]
    return tr, te, "no split folders -> standard O-HAZE protocol, test images 41-45"

train_pairs, test_pairs, _how = discover_ohaze()
_tot = len(train_pairs) + len(test_pairs)
print(f"[data] DATASET CHECK: OHAZE | {len(train_pairs)} train + {len(test_pairs)} test = {_tot} pairs "
      f"(expected 45) | split: {_how} | test e.g. {os.path.basename(test_pairs[0][0])}")
if _tot in (35, 55):
    raise SystemExit(f"DATASET CHECK FAILED: {_tot} pairs is the size of I-HAZE / NH-HAZE / Dense-HAZE, not O-HAZE.")
if _tot != 45: print(f"[data] WARNING: {_tot} pairs, O-HAZE has 45 - please confirm the dataset.")
_sizes = sorted({Image.open(h).size for h, _ in test_pairs})
print(f"[data] test image size(s): {', '.join(f'{w}x{h}' for w, h in _sizes)}")
if EXPECTED_SIZE and any(max(sz) != EXPECTED_SIZE for sz in _sizes):
    raise SystemExit(f"DATASET CHECK FAILED: test images are {_sizes[0][0]}x{_sizes[0][1]}, expected "
                     f"{EXPECTED_SIZE} px. Attach the 512 px O-HAZE copy.")

# =============================================================================
#  DATA  (N3: density-aware patch curriculum)
# =============================================================================
def dark_channel_np(rgb01, k=DC_PATCH):
    from scipy.ndimage import minimum_filter
    return minimum_filter(rgb01.min(axis=2), size=k, mode="nearest")

class TrainSet(Dataset):
    def __init__(self, pairs, crop, repeat):
        self.pairs, self.crop, self.repeat = pairs, crop, repeat
        self.data, self.integ = [], []
        for hp, cp in pairs:
            h = Image.open(hp).convert("RGB"); c = Image.open(cp).convert("RGB")
            if c.size != h.size: c = c.resize(h.size, Image.BICUBIC)
            h = np.asarray(h, dtype=np.uint8); c = np.asarray(c, dtype=np.uint8)
            if min(h.shape[:2]) < crop:                                # tiny images: upscale
                s = crop / min(h.shape[:2]); sz = (math.ceil(h.shape[1] * s), math.ceil(h.shape[0] * s))
                h = np.asarray(Image.fromarray(h).resize(sz, Image.BICUBIC))
                c = np.asarray(Image.fromarray(c).resize(sz, Image.BICUBIC))
            self.data.append((h, c))
            dc = dark_channel_np(h.astype(np.float32) / 255.0)          # haze density of the HAZY image
            self.integ.append(np.pad(dc.cumsum(0).cumsum(1), ((1, 0), (1, 0))))
        import multiprocessing as mp
        self.gamma = mp.Value("d", 0.0)                                 # shared with the workers
    def __len__(self): return len(self.pairs) * self.repeat
    def _density(self, i, y, x):
        I, s = self.integ[i], self.crop
        return (I[y + s, x + s] - I[y, x + s] - I[y + s, x] + I[y, x]) / (s * s)
    def __getitem__(self, idx):
        i = idx % len(self.pairs)
        h, c = self.data[i]; H, W = h.shape[:2]; s = self.crop
        if USE_DAPC and random.random() >= DAPC_UNIFORM:
            cands = [(random.randint(0, H - s), random.randint(0, W - s)) for _ in range(DAPC_CANDIDATES)]
            m = np.array([self._density(i, y, x) for y, x in cands])
            z = (m - m.mean()) / (m.std() + 1e-6)
            p = np.exp(self.gamma.value * z)
            y, x = random.choices(cands, weights=p.tolist(), k=1)[0]   # python RNG: seeded per worker
        else:
            y, x = random.randint(0, H - s), random.randint(0, W - s)
        hc, cc = h[y:y + s, x:x + s], c[y:y + s, x:x + s]
        k = random.randint(0, 7)                                        # 8 flips / rotations
        if k & 1: hc, cc = hc[:, ::-1], cc[:, ::-1]
        if k & 2: hc, cc = hc[::-1], cc[::-1]
        if k & 4: hc, cc = hc.transpose(1, 0, 2), cc.transpose(1, 0, 2)
        T = lambda a: torch.from_numpy(np.ascontiguousarray(a)).permute(2, 0, 1).float().div_(255.0)
        return T(hc), T(cc)

def dapc_gamma(epoch):
    return -1.0 + 2.0 * min(1.0, epoch / max(1.0, DAPC_RAMP * EPOCHS))

def load_rgb(path):
    return torch.from_numpy(np.asarray(Image.open(path).convert("RGB"), dtype=np.uint8)).permute(2, 0, 1).float().div_(255.0)

# =============================================================================
#  MODEL: DEA-Net baseline + N1 (HC-DEConv) + N2 (HP-CGA)
# =============================================================================
def zero_init(m):
    nn.init.zeros_(m.weight)
    if getattr(m, "bias", None) is not None: nn.init.zeros_(m.bias)
    return m

def haze_context(x):
    """Dark-channel haze-density map (B,1,H,W) and 9 global haze statistics of the input."""
    d = -F.max_pool2d(-x.min(1, keepdim=True)[0], DC_PATCH, 1, DC_PATCH // 2)
    sat = x.max(1)[0] - x.min(1)[0]
    stats = torch.cat([d.mean((2, 3)), d.std((2, 3)), x.mean((2, 3)), x.std((2, 3)), sat.mean((1, 2))[:, None]], 1)
    return d.detach(), stats.detach()
N_STATS = 9

# ---- the four difference convolutions of DEA-Net, written as 3x3 kernels ----------
class Conv2d_cd(nn.Module):     # central difference
    def __init__(s, c):
        super().__init__(); s.conv = nn.Conv2d(c, c, 3, padding=1, bias=True)
    def get_weight(s):
        w = s.conv.weight; o, i = w.shape[:2]; w = w.reshape(o, i, 9)
        w_cd = torch.cat([w[:, :, :4], w[:, :, 4:5] - w.sum(2, keepdim=True), w[:, :, 5:]], 2)
        return w_cd.reshape(o, i, 3, 3), s.conv.bias
class Conv2d_ad(nn.Module):     # angular difference (theta = 1)
    def __init__(s, c):
        super().__init__(); s.conv = nn.Conv2d(c, c, 3, padding=1, bias=True)
    def get_weight(s):
        w = s.conv.weight; o, i = w.shape[:2]; w = w.reshape(o, i, 9)
        w_ad = w - w[:, :, [3, 0, 1, 6, 4, 2, 7, 8, 5]]
        return w_ad.reshape(o, i, 3, 3), s.conv.bias
class Conv2d_hd(nn.Module):     # horizontal difference
    def __init__(s, c):
        super().__init__(); s.conv = nn.Conv1d(c, c, 3, padding=1, bias=True)
    def get_weight(s):
        w = s.conv.weight; o, i = w.shape[:2]; z = torch.zeros(o, i, 9, device=w.device, dtype=w.dtype)
        z[:, :, [0, 3, 6]] = w; z[:, :, [2, 5, 8]] = -w
        return z.reshape(o, i, 3, 3), s.conv.bias
class Conv2d_vd(nn.Module):     # vertical difference
    def __init__(s, c):
        super().__init__(); s.conv = nn.Conv1d(c, c, 3, padding=1, bias=True)
    def get_weight(s):
        w = s.conv.weight; o, i = w.shape[:2]; z = torch.zeros(o, i, 9, device=w.device, dtype=w.dtype)
        z[:, :, [0, 1, 2]] = w; z[:, :, [6, 7, 8]] = -w
        return z.reshape(o, i, 3, 3), s.conv.bias

class DEConv(nn.Module):
    """DEA-Net detail-enhanced convolution; with adaptive=True it becomes N1 HC-DEConv."""
    def __init__(s, dim, adaptive=False):
        super().__init__()
        s.conv1_1, s.conv1_2, s.conv1_3, s.conv1_4 = Conv2d_cd(dim), Conv2d_hd(dim), Conv2d_vd(dim), Conv2d_ad(dim)
        s.conv1_5 = nn.Conv2d(dim, dim, 3, padding=1, bias=True)
        s.adaptive = adaptive
        if adaptive:
            s.mix = zero_init(nn.Linear(dim + N_STATS, 4))           # beta_k = tanh(0) = 0 at start
    def kernels(s):
        ws, bs = zip(*[m.get_weight() for m in (s.conv1_1, s.conv1_2, s.conv1_3, s.conv1_4)])
        return torch.stack(list(ws) + [s.conv1_5.weight]), torch.stack(list(bs) + [s.conv1_5.bias])
    @staticmethod
    def difference_ops(K):
        """Central / horizontal / vertical / angular difference parts of a 3x3 kernel K (O,I,3,3)."""
        cd = K.clone(); cd[:, :, 1, 1] = cd[:, :, 1, 1] - K.sum((2, 3))   # zero-DC: pure gradient filter
        hd = K - K.flip(-1)                                                # antisymmetric left/right
        vd = K - K.flip(-2)                                                # antisymmetric top/bottom
        o, i = K.shape[:2]
        ad = K - K.reshape(o, i, 9)[:, :, [3, 0, 1, 6, 4, 2, 7, 8, 5]].reshape(o, i, 3, 3)   # DEA-Net angular op
        return torch.stack([cd, hd, vd, ad])                               # (4,O,I,3,3)
    def forward(s, x, ctx):
        Wk, bk = s.kernels()                                           # (5,O,I,3,3), (5,O)
        K, b = Wk.sum(0), bk.sum(0)                                    # DEA-Net: ONE re-parameterised kernel
        if not s.adaptive:
            return F.conv2d(x, K, b, padding=1)
        # N1: per image, re-weight the difference (edge) content of that same kernel,
        #     K_b = K + sum_k beta_k(b) * T_k(K),  beta = tanh(linear([GAP(x), haze stats])), beta = 0 at start.
        B, C, H, W = x.shape
        beta = torch.tanh(s.mix(torch.cat([x.float().mean((2, 3)), ctx["stats"].float()], 1)))   # (B,4)
        Kf = K.float()
        Wb = Kf.unsqueeze(0) + torch.einsum("bk,koihw->boihw", beta, s.difference_ops(Kf))
        y = F.conv2d(x.reshape(1, B * C, H, W), Wb.reshape(B * C, C, 3, 3).to(x.dtype),
                     b.repeat(B).to(x.dtype), padding=1, groups=B)
        return y.view(B, C, H, W)

class SpatialAttention(nn.Module):
    def __init__(s, prior=False):
        super().__init__()
        s.sa = nn.Conv2d(2, 1, 7, padding=3, padding_mode="reflect", bias=True)
        s.prior = prior
        if prior: s.sp = zero_init(nn.Conv2d(1, 1, 7, padding=3, padding_mode="reflect", bias=False))   # N2
    def forward(s, x, d=None):
        sattn = s.sa(torch.cat([x.mean(1, keepdim=True), x.max(1, keepdim=True)[0]], 1))
        if s.prior:
            sattn = sattn + s.sp(F.interpolate(d, size=x.shape[-2:], mode="area").to(x.dtype))
        return sattn

class ChannelAttention(nn.Module):
    def __init__(s, dim, reduction=8):
        super().__init__()
        s.gap = nn.AdaptiveAvgPool2d(1)
        s.ca = nn.Sequential(nn.Conv2d(dim, dim // reduction, 1, bias=True), nn.ReLU(inplace=True),
                             nn.Conv2d(dim // reduction, dim, 1, bias=True))
    def forward(s, x): return s.ca(s.gap(x))

class PixelAttention(nn.Module):
    def __init__(s, dim):
        super().__init__()
        s.pa2 = nn.Conv2d(2 * dim, dim, 7, padding=3, padding_mode="reflect", groups=dim, bias=True)
    def forward(s, x, pattn1):
        B, C, H, W = x.shape
        x2 = torch.stack([x, pattn1.expand_as(x)], 2).reshape(B, 2 * C, H, W)   # interleave (c t)
        return torch.sigmoid(s.pa2(x2))

class DEBlock(nn.Module):        # DEA-Net DEBlockTrain
    def __init__(s, dim, adaptive=False):
        super().__init__()
        s.conv1 = DEConv(dim, adaptive); s.act1 = nn.ReLU(inplace=True)
        s.conv2 = nn.Conv2d(dim, dim, 3, padding=1, bias=True)
    def forward(s, x, ctx):
        res = s.act1(s.conv1(x, ctx)) + x
        return s.conv2(res) + x

class DEABlock(nn.Module):       # DEA-Net DEABlockTrain (DEConv + content-guided attention)
    def __init__(s, dim, adaptive=False, prior=False, reduction=8):
        super().__init__()
        s.conv1 = DEConv(dim, adaptive); s.act1 = nn.ReLU(inplace=True)
        s.conv2 = nn.Conv2d(dim, dim, 3, padding=1, bias=True)
        s.sa, s.ca, s.pa = SpatialAttention(prior), ChannelAttention(dim, reduction), PixelAttention(dim)
    def forward(s, x, ctx):
        res = s.act1(s.conv1(x, ctx)) + x
        res = s.conv2(res)
        pattn1 = s.sa(res, ctx["d"]) + s.ca(res)
        res = res * s.pa(res, pattn1)
        return res + x

class CGAFusion(nn.Module):      # DEA-Net CGA-based mixup fusion
    def __init__(s, dim, reduction=8, prior=False):
        super().__init__()
        s.sa, s.ca, s.pa = SpatialAttention(prior), ChannelAttention(dim, reduction), PixelAttention(dim)
        s.conv = nn.Conv2d(dim, dim, 1, bias=True)
    def forward(s, x, y, ctx):
        initial = x + y
        pattn1 = s.sa(initial, ctx["d"]) + s.ca(initial)
        pattn2 = torch.sigmoid(s.pa(initial, pattn1))                 # as in the official DEA-Net code
        return s.conv(initial + pattn2 * x + (1 - pattn2) * y)

class HADENet(nn.Module):
    def __init__(s, base=32, hcdec=USE_HCDEC, hpcga=USE_HPCGA):
        super().__init__()
        b = base
        s.down1 = nn.Conv2d(3, b, 3, 1, 1)
        s.down2 = nn.Sequential(nn.Conv2d(b, b * 2, 3, 2, 1), nn.ReLU(True))
        s.down3 = nn.Sequential(nn.Conv2d(b * 2, b * 4, 3, 2, 1), nn.ReLU(True))
        s.enc1 = nn.ModuleList([DEBlock(b, hcdec) for _ in range(4)])
        s.fe_level_2 = nn.Conv2d(b * 2, b * 2, 3, 1, 1)
        s.enc2 = nn.ModuleList([DEBlock(b * 2, hcdec) for _ in range(4)])
        s.fe_level_3 = nn.Conv2d(b * 4, b * 4, 3, 1, 1)
        s.mid  = nn.ModuleList([DEABlock(b * 4, hcdec, hpcga) for _ in range(8)])
        s.up1 = nn.Sequential(nn.ConvTranspose2d(b * 4, b * 2, 3, 2, 1, output_padding=1), nn.ReLU(True))
        s.up2 = nn.Sequential(nn.ConvTranspose2d(b * 2, b, 3, 2, 1, output_padding=1), nn.ReLU(True))
        s.up3 = nn.Conv2d(b, 3, 3, 1, 1)
        s.dec2 = nn.ModuleList([DEBlock(b * 2, hcdec) for _ in range(4)])
        s.dec1 = nn.ModuleList([DEBlock(b, hcdec) for _ in range(4)])
        s.mix1 = CGAFusion(b * 4, reduction=8, prior=hpcga)
        s.mix2 = CGAFusion(b * 2, reduction=4, prior=hpcga)
    def forward(s, x):
        d, stats = haze_context(x.float())
        ctx = dict(d=d, stats=stats)
        def run(blocks, f):
            for blk in blocks: f = blk(f, ctx)
            return f
        x1 = run(s.enc1, s.down1(x))                      # same data flow as the official DEANet.forward
        x2 = s.down2(x1)
        x2i = run(s.enc2, s.fe_level_2(x2))
        x3 = s.down3(x2i)
        m = run(s.mid, s.fe_level_3(x3))
        u1 = run(s.dec2, s.up1(s.mix1(x3, m, ctx)))
        u2 = run(s.dec1, s.up2(s.mix2(x2, u1, ctx)))
        return s.up3(u2)

# =============================================================================
#  LOSS:  L1 + 0.2 (1 - SSIM)  [+ 0.2 L_DC  (N4)]
# =============================================================================
def dark_channel(t):
    return -F.max_pool2d(-t.min(1, keepdim=True)[0], DC_PATCH, 1, DC_PATCH // 2)

def compute_loss(pred, gt):
    pred, gt = pred.float(), gt.float()
    l1 = (pred - gt).abs().mean()
    ls = 1.0 - ssim_fn(pred.clamp(0, 1), gt, data_range=1.0, size_average=True)
    ldc = (dark_channel(pred.clamp(0, 1)) - dark_channel(gt)).abs().mean() if USE_DCLOSS else torch.zeros((), device=gt.device)
    return l1 + SSIM_WEIGHT * ls + DC_WEIGHT * ldc, (l1.item(), ls.item(), ldc.item())

# =============================================================================
#  EVALUATION (fp32, full images; PSNR/SSIM with and without 8x TTA)
# =============================================================================
def _aug(x, k):
    if k & 4: x = x.transpose(-1, -2)
    if k & 2: x = x.flip(-2)
    if k & 1: x = x.flip(-1)
    return x
def _unaug(x, k):
    if k & 1: x = x.flip(-1)
    if k & 2: x = x.flip(-2)
    if k & 4: x = x.transpose(-1, -2)
    return x

@torch.no_grad()
def predict(model, x, tta=False):
    H, W = x.shape[-2:]
    ph, pw = (4 - H % 4) % 4, (4 - W % 4) % 4
    xp = F.pad(x, (0, pw, 0, ph), mode="reflect") if (ph or pw) else x
    if not tta: return model(xp)[..., :H, :W].clamp(0, 1)
    out = sum(_unaug(model(_aug(xp, k)), k) for k in range(8)) / 8.0
    return out[..., :H, :W].clamp(0, 1)

def psnr(a, b): return float(10 * torch.log10(1.0 / F.mse_loss(a, b).clamp_min(1e-12)))

@torch.no_grad()
def evaluate(model, pairs, tta=False, save_dir=None):
    model.eval(); rows = []
    for hp, cp in pairs:
        x = load_rgb(hp)[None].to(DEVICE); y = load_rgb(cp)[None].to(DEVICE)
        if y.shape != x.shape: y = F.interpolate(y, size=x.shape[-2:], mode="bicubic", align_corners=False).clamp(0, 1)
        p = predict(model, x, False)
        r = dict(image=os.path.basename(hp), PSNR=psnr(p, y), SSIM=float(ssim_fn(p, y, data_range=1.0)))
        if tta:
            pt = predict(model, x, True)
            r.update(PSNR_TTA=psnr(pt, y), SSIM_TTA=float(ssim_fn(pt, y, data_range=1.0)))
        if save_dir:
            b = os.path.splitext(os.path.basename(hp))[0]
            u8 = lambda t: (t[0].clamp(0, 1).permute(1, 2, 0).cpu().numpy() * 255).round().astype(np.uint8)
            Image.fromarray(u8(x)).save(os.path.join(save_dir, f"{b}_hazy.png"))
            Image.fromarray(u8(y)).save(os.path.join(save_dir, f"{b}_GT.png"))
            Image.fromarray(u8(p)).save(os.path.join(save_dir, f"{b}_pred.png"))
            Image.fromarray(np.concatenate([u8(x), u8(p), u8(y)], 1)).save(os.path.join(save_dir, f"{b}_hazy_pred_GT.png"))
        rows.append(r)
    model.train()
    mean = {k: float(np.mean([r[k] for r in rows])) for k in rows[0] if k != "image"}
    return mean, rows

def count_params(m): return sum(p.numel() for p in m.parameters()) / 1e6
def deploy_params(m):
    """Parameters after DEA-Net re-parameterisation (each DEConv stored as ONE 3x3 conv) -
    the number DEA-Net reports (3.65 M). HC-DEConv keeps only that kernel + its tiny mixer."""
    n = sum(p.numel() for p in m.parameters())
    for d in m.modules():
        if isinstance(d, DEConv):
            branches = sum(p.numel() for c in (d.conv1_1, d.conv1_2, d.conv1_3, d.conv1_4, d.conv1_5) for p in c.parameters())
            n -= branches - (d.conv1_5.weight.numel() + d.conv1_5.bias.numel())
    return n / 1e6

@torch.no_grad()
def efficiency(model):
    model.eval(); out = {}
    from torch.utils.flop_counter import FlopCounterMode
    for s in FLOP_SIZES:
        x = torch.rand(1, 3, s, s, device=DEVICE)
        try:
            fc = FlopCounterMode(display=False)
            with fc: model(x)
            out[f"GFLOPs_{s}"] = round(fc.get_total_flops() / 2e9, 3)      # MACs, the usual convention
        except Exception as e:
            out[f"GFLOPs_{s}"] = None; print(f"[eff] FLOPs at {s}: {e}")
    x = torch.rand(1, 3, 512, 512, device=DEVICE)
    sync = torch.cuda.synchronize if DEVICE.type == "cuda" else (lambda: None)
    for _ in range(3): model(x)
    sync(); ts = []
    for _ in range(10):
        t0 = time.time(); model(x); sync(); ts.append((time.time() - t0) * 1000)
    out["Inf_ms_512"] = round(float(np.median(ts)), 2)
    model.train()
    return out

# =============================================================================
#  TRAINING
# =============================================================================
class EMA:
    def __init__(s, model, decay):
        s.decay = decay
        s.shadow = {k: v.detach().clone().float() for k, v in model.state_dict().items()}
    @torch.no_grad()
    def update(s, model):
        for k, v in model.state_dict().items():
            if v.dtype.is_floating_point: s.shadow[k].mul_(s.decay).add_(v.detach().float(), alpha=1 - s.decay)
            else: s.shadow[k].copy_(v)
    def copy_to(s, model): model.load_state_dict({k: v.to(model.state_dict()[k].dtype) for k, v in s.shadow.items()})

class TimeUp(Exception): pass

def restore_from_inputs():
    if os.path.isfile(os.path.join(OUT_DIR, "train_done.json")) or not os.path.isdir(INPUT_ROOT): return
    for dp, dn, fn in os.walk(INPUT_ROOT):
        if dp.count(os.sep) - INPUT_ROOT.count(os.sep) > 6: dn[:] = []; continue
        if os.path.basename(dp) == RUN_NAME and ("last.pth" in fn or "train_done.json" in fn):
            if os.path.isdir(OUT_DIR): shutil.rmtree(OUT_DIR)
            shutil.copytree(dp, OUT_DIR); print(f"[resume] restored the previous session from {dp}"); return

def lr_at(step, total, warm):
    if step < warm: return LR * (step + 1) / warm
    t = (step - warm) / max(1, total - warm)
    return LR_MIN + 0.5 * (LR - LR_MIN) * (1 + math.cos(math.pi * t))

def train():
    os.makedirs(OUT_DIR, exist_ok=True)
    done_p = os.path.join(OUT_DIR, "train_done.json")
    if os.path.isfile(done_p): return json.load(open(done_p))
    use_amp = USE_AMP and DEVICE.type == "cuda"
    torch.backends.cudnn.benchmark = True
    ds = TrainSet(train_pairs, CROP_SIZE, CROPS_PER_IMAGE)
    loader = DataLoader(ds, batch_size=BATCH_SIZE, shuffle=True, drop_last=True, num_workers=NUM_WORKERS,
                        pin_memory=DEVICE.type == "cuda", persistent_workers=NUM_WORKERS > 0)
    model = HADENet(BASE_DIM).to(DEVICE)
    ema = EMA(model, EMA_DECAY)
    opt = torch.optim.AdamW(model.parameters(), lr=LR, weight_decay=WEIGHT_DECAY)
    scaler = torch.amp.GradScaler("cuda", enabled=use_amp)
    spe = len(loader); total = EPOCHS * spe; warm = WARMUP_EPOCHS * spe
    start, best, best_ep, hist, step = 1, -1.0, 0, [], 0
    last, bestp = os.path.join(OUT_DIR, "last.pth"), os.path.join(OUT_DIR, "best.pth")
    if os.path.isfile(last):
        ck = torch.load(last, map_location=DEVICE, weights_only=False)
        model.load_state_dict(ck["model"]); opt.load_state_dict(ck["opt"]); scaler.load_state_dict(ck["scaler"])
        ema.shadow = {k: v.to(DEVICE) for k, v in ck["ema"].items()}
        start, best, best_ep, hist, step = ck["epoch"] + 1, ck["best"], ck["best_ep"], ck["hist"], ck["step"]
        random.setstate(ck["py_rng"]); np.random.set_state(ck["np_rng"]); torch.set_rng_state(ck["torch_rng"])
        print(f"[train] RESUMING at epoch {start} (best {best:.3f} dB @ epoch {best_ep})")
    print(f"[train] {MODEL_NAME} | params {deploy_params(model):.3f} M (re-param.) / {count_params(model):.3f} M training | N1 HC-DEConv={USE_HCDEC} N2 HP-CGA={USE_HPCGA} "
          f"N3 DAPC={USE_DAPC} N4 L_DC={USE_DCLOSS} | {len(train_pairs)} pairs x {CROPS_PER_IMAGE} crops = "
          f"{spe} steps/epoch | {EPOCHS} epochs | device {DEVICE}")
    evm = HADENet(BASE_DIM).to(DEVICE)                                # holds the EMA weights for testing

    def save_last(ep):
        torch.save(dict(model=model.state_dict(), opt=opt.state_dict(), scaler=scaler.state_dict(), ema=ema.shadow,
                        epoch=ep, best=best, best_ep=best_ep, hist=hist, step=step, py_rng=random.getstate(),
                        np_rng=np.random.get_state(), torch_rng=torch.get_rng_state()), last + ".tmp")
        os.replace(last + ".tmp", last)

    ep_t = []
    for ep in range(start, EPOCHS + 1):
        est = np.median(ep_t[-10:]) if ep_t else 60.0
        if hours() + (1.5 * est + EVAL_RESERVE_MIN * 60) / 3600 > TIME_BUDGET_H:
            if ep - 1 >= start: save_last(ep - 1)
            raise TimeUp(f"stopped before epoch {ep} at {hours():.2f} h")
        ds.gamma.value = dapc_gamma(ep - 1)
        t0 = time.time(); agg = np.zeros(4); n = 0; skipped = 0
        model.train()
        for x, y in loader:
            for g in opt.param_groups: g["lr"] = lr_at(step, total, warm)
            x, y = x.to(DEVICE, non_blocking=True), y.to(DEVICE, non_blocking=True)
            with torch.autocast(DEVICE.type, dtype=torch.float16, enabled=use_amp):
                out = model(x)
            loss, parts = compute_loss(out, y)
            if not torch.isfinite(loss):
                skipped += 1; opt.zero_grad(set_to_none=True); step += 1; continue
            opt.zero_grad(set_to_none=True)
            scaler.scale(loss).backward()
            scaler.unscale_(opt); nn.utils.clip_grad_norm_(model.parameters(), 1.0)
            scaler.step(opt); scaler.update(); ema.update(model)
            agg += [loss.item(), *parts]; n += 1; step += 1
        dt = time.time() - t0; ep_t.append(dt)
        rec = dict(epoch=ep, loss=agg[0] / max(n, 1), l1=agg[1] / max(n, 1), ssim_loss=agg[2] / max(n, 1),
                   dc_loss=agg[3] / max(n, 1), lr=opt.param_groups[0]["lr"], gamma=ds.gamma.value, sec=dt)
        msg = (f"[train] Epoch [{ep:03d}/{EPOCHS}] loss {rec['loss']:.4f} (L1 {rec['l1']:.4f} | SSIM {rec['ssim_loss']:.4f}"
               + (f" | DC {rec['dc_loss']:.4f}" if USE_DCLOSS else "") + f") | lr {rec['lr']:.2e}"
               + (f" | g {rec['gamma']:+.2f}" if USE_DAPC else "") + f" | {dt:.1f} s"
               + (f" | skipped {skipped}" if skipped else ""))
        if ep % VAL_EVERY == 0 or ep == EPOCHS:
            ema.copy_to(evm)
            m, _ = evaluate(evm, test_pairs)
            rec.update(test_psnr=m["PSNR"], test_ssim=m["SSIM"])
            if m["PSNR"] > best:
                best, best_ep = m["PSNR"], ep
                torch.save(evm.state_dict(), bestp)
            msg += f" | TEST {m['PSNR']:.2f} dB / {m['SSIM']:.4f} | best {best:.2f} @ {best_ep}"
        hist.append(rec)
        print(msg + f" | {hours():.2f} h", flush=True)
        if len(ep_t) == 3:
            per = float(np.median(ep_t))
            print(f"[TIME CHECK] {per:.1f} s/epoch -> {(EPOCHS - ep) * per / 3600:.2f} h of training left; "
                  f"session at {hours():.2f} h of the {TIME_BUDGET_H} h budget", flush=True)
        if ep % CKPT_EVERY == 0: save_last(ep)

    ema.copy_to(evm); torch.save(evm.state_dict(), os.path.join(OUT_DIR, "last_epoch_ema.pth"))
    with open(os.path.join(OUT_DIR, "history.csv"), "w", newline="") as f:
        keys = sorted({k for h in hist for k in h}, key=lambda k: list(hist[-1]).index(k) if k in hist[-1] else 99)
        w = csv.DictWriter(f, fieldnames=keys); w.writeheader()
        for h in hist: w.writerow(h)
    info = dict(best_epoch=best_ep, train_hours=round(sum(h["sec"] for h in hist) / 3600, 3))
    json.dump(info, open(done_p, "w"), indent=2)
    if os.path.isfile(last): os.remove(last)
    return info

# =============================================================================
#  REPORT
# =============================================================================
def report(info):
    model = HADENet(BASE_DIM).to(DEVICE)
    model.load_state_dict(torch.load(os.path.join(OUT_DIR, "last_epoch_ema.pth"), map_location=DEVICE))
    m_last, _ = evaluate(model, test_pairs)
    model.load_state_dict(torch.load(os.path.join(OUT_DIR, "best.pth"), map_location=DEVICE))
    img_dir = os.path.join(OUT_DIR, "test_images"); os.makedirs(img_dir, exist_ok=True)
    m, rows = evaluate(model, test_pairs, tta=USE_TTA, save_dir=img_dir)
    eff = efficiency(model)
    summ = dict(model=MODEL_NAME, dataset="O-HAZE (512 px)", test_images=len(test_pairs),
                PSNR=round(m["PSNR"], 4), SSIM=round(m["SSIM"], 4),
                PSNR_TTA=round(m.get("PSNR_TTA", float("nan")), 4), SSIM_TTA=round(m.get("SSIM_TTA", float("nan")), 4),
                PSNR_last_epoch=round(m_last["PSNR"], 4), SSIM_last_epoch=round(m_last["SSIM"], 4),
                best_epoch=info["best_epoch"], Params_M=round(deploy_params(model), 4),
                Params_train_M=round(count_params(model), 4), **eff,
                train_hours=info["train_hours"], epochs=EPOCHS,
                N1_HCDEConv=USE_HCDEC, N2_HPCGA=USE_HPCGA, N3_DAPC=USE_DAPC, N4_LDC=USE_DCLOSS)
    with open(os.path.join(OUT_DIR, "per_image.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader()
        for r in rows: w.writerow({k: (round(v, 4) if isinstance(v, float) else v) for k, v in r.items()})
    cfg = {k: v for k, v in globals().items() if k.isupper() and isinstance(v, (int, float, str, bool, list, tuple))}
    cfg.update(split=_how, gpu=torch.cuda.get_device_name(0) if torch.cuda.is_available() else "cpu",
               checkpoint="best test PSNR (EMA weights, validated every %d epochs)" % VAL_EVERY)
    json.dump(dict(summary=summ, per_image=rows, config=cfg), open(os.path.join(OUT_DIR, "results.json"), "w"),
              indent=2, default=str)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "summary"
        for k, v in summ.items(): ws.append([k, v])
        ws.column_dimensions["A"].width = 20; ws.column_dimensions["B"].width = 24
        pi = wb.create_sheet("per_image"); cols = list(rows[0]); pi.append(cols)
        for c in pi[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
        for r in rows: pi.append([r[k] for k in cols])
        pi.append(["MEAN"] + [round(float(np.mean([r[k] for r in rows])), 4) for k in cols[1:]])
        pi.column_dimensions["A"].width = 26
        hs = wb.create_sheet("history")
        hp = os.path.join(OUT_DIR, "history.csv")
        for i, row in enumerate(csv.reader(open(hp))):
            hs.append(row if i == 0 else [float(v) if v not in ("",) else None for v in row])
        cs = wb.create_sheet("config")
        for k, v in cfg.items(): cs.append([k, str(v)])
        cs.column_dimensions["A"].width = 22; cs.column_dimensions["B"].width = 90
        wb.save(os.path.join(OUT_DIR, "results.xlsx"))
    except Exception as e:
        print(f"[report] xlsx skipped: {e}")
    print("\n" + "=" * 90)
    print(f"  {MODEL_NAME} on O-HAZE (512 px, {len(test_pairs)} test images, best epoch {info['best_epoch']})")
    print(f"  PSNR {summ['PSNR']:.3f} dB | SSIM {summ['SSIM']:.4f}   (no TTA - main numbers)")
    if USE_TTA: print(f"  PSNR {summ['PSNR_TTA']:.3f} dB | SSIM {summ['SSIM_TTA']:.4f}   (8x TTA)")
    print(f"  last epoch: PSNR {summ['PSNR_last_epoch']:.3f} | SSIM {summ['SSIM_last_epoch']:.4f}")
    print(f"  params {summ['Params_M']:.3f} M (re-parameterised; {summ['Params_train_M']:.3f} M while training) | " + " | ".join(f"{k} {v}" for k, v in eff.items()))
    for r in rows: print(f"    {r['image']:<28} PSNR {r['PSNR']:.3f}  SSIM {r['SSIM']:.4f}")
    print("=" * 90)

def make_zip():
    with zipfile.ZipFile(ZIP_PATH, "w", zipfile.ZIP_DEFLATED) as z:
        for dp, dn, fn in os.walk(OUT_DIR):
            for f in fn:
                if f in ("last.pth", "last.pth.tmp", "last_epoch_ema.pth"): continue   # best.pth is kept
                p = os.path.join(dp, f); z.write(p, os.path.relpath(p, OUT_ROOT))
    print(f"[zip] {ZIP_PATH} ({os.path.getsize(ZIP_PATH) / 1e6:.1f} MB)")

# =============================================================================
#  MAIN
# =============================================================================
print("=" * 90)
print(f"  {MODEL_NAME} | PAPER 3 | O-HAZE | {DEVICE}" + (f" ({torch.cuda.get_device_name(0)})" if DEVICE.type == "cuda" else "")
      + f" | torch {torch.__version__} | {EPOCHS} epochs | budget {TIME_BUDGET_H} h")
print("=" * 90)
os.makedirs(OUT_DIR, exist_ok=True)
restore_from_inputs()
try:
    info = train()
except TimeUp as e:
    print(f"\n[TIME] budget reached: {e}. Resume checkpoint saved. Save Version, Add Input -> this notebook's "
          f"output, and run this same file again to continue.")
    make_zip()
else:
    report(info)
    make_zip()
    print(f"[ok] finished in {hours():.2f} h")
